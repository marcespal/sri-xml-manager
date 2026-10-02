"""
excel_writer.py — Genera reportes Excel: FC_Recibidas y CR_Recibidas.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from xml_parser import Factura, Retencion
from utils import get_logger, timestamp

logger = get_logger("excel_writer")

COLOR_H_FC  = "1F4E79"
COLOR_H_CR  = "833C00"
COLOR_TXT   = "FFFFFF"
COLOR_PAR   = "EBF3FB"
COLOR_IMPAR = "FFFFFF"
COLOR_TOTAL = "D6E4F0"


def _h(color_bg: str = COLOR_H_FC):
    return dict(
        font=Font(bold=True, color=COLOR_TXT, size=10),
        fill=PatternFill("solid", fgColor=color_bg),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
    )


def _ap(c, **kw):
    for k, v in kw.items():
        setattr(c, k, v)


COLS_FC = [
    ("TIPO", 12), ("FECHA EMISION", 16), ("RUC EMISOR", 15), ("RAZON SOCIAL EMISOR", 35),
    ("NUMERO", 18), ("RUC/CI COMPRADOR", 15), ("COMPRADOR", 30),
    ("BASE 0%", 12), ("BASE 12%", 12), ("BASE 14%", 12), ("BASE 15%", 12),
    ("BASE NO OBJETO", 14), ("BASE EXENTA", 12), ("SUBTOTAL", 12),
    ("DESCUENTO", 12), ("IVA", 12), ("PROPINA", 10), ("TOTAL", 12),
    ("CLAVE ACCESO", 52),
]
# Columnas monetarias desplazadas +1 por la nueva columna TIPO
MONED_FC = [8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18]

COLS_CR = [
    ("FECHA EMISION", 16), ("RUC RETENEDOR", 15), ("RAZON RETENEDOR", 35),
    ("NUMERO RETENCION", 20), ("RUC/CI RETENIDO", 15), ("RAZON RETENIDO", 30),
    ("PERIODO FISCAL", 14), ("COD DOC SUSTENTO", 18), ("NUM DOC SUSTENTO", 22),
    ("FECHA DOC SUSTENTO", 18), ("TIPO IMPUESTO", 14), ("COD RETENCION", 14),
    ("BASE IMPONIBLE", 14), ("% RETENCION", 12), ("VALOR RETENIDO", 14),
    ("CLAVE ACCESO", 52),
]
MONED_CR = [13, 14, 15]


def _encabezados(ws, cols, color):
    est = _h(color)
    for col, (nombre, ancho) in enumerate(cols, 1):
        c = ws.cell(row=1, column=col, value=nombre)
        _ap(c, **est)
        ws.column_dimensions[get_column_letter(col)].width = ancho


def _fila_moneda(ws, fila, cols_moneda):
    for col in cols_moneda:
        ws.cell(row=fila, column=col).number_format = '#,##0.00'


def _fila_total(ws, fila_total, fila_ini, cols_moneda, color):
    fill = PatternFill("solid", fgColor=color)
    ws.cell(row=fila_total, column=1, value="TOTAL").font = Font(bold=True)
    for col in cols_moneda:
        cl = get_column_letter(col)
        c = ws.cell(row=fila_total, column=col,
                    value=f"=SUM({cl}{fila_ini}:{cl}{fila_total - 1})")
        c.font = Font(bold=True)
        c.fill = fill
        c.number_format = '#,##0.00'


def escribir_facturas(ws, facturas: list):
    ws.title = "FC_Recibidas"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 35
    _encabezados(ws, COLS_FC, COLOR_H_FC)

    TIPO_NOM = {"01": "FACTURA", "03": "LIQUIDACION", "04": "N. CREDITO",
                "05": "N. DEBITO", "07": "RETENCION"}
    for fila, f in enumerate(facturas, 2):
        fill = PatternFill("solid", fgColor=COLOR_PAR if fila % 2 == 0 else COLOR_IMPAR)
        numero = f"{f.establecimiento}-{f.punto_emision}-{f.secuencial}"
        tipo_nom = TIPO_NOM.get(getattr(f, "cod_doc", "01"), "FACTURA")
        vals = [
            tipo_nom, f.fecha_emision, f.ruc, f.razon_social, numero,
            f.identificacion_comprador, f.razon_social_comprador,
            f.base_iva_0, f.base_iva_12, f.base_iva_14, f.base_iva_15,
            f.base_no_objeto, f.base_exenta, f.total_sin_impuestos,
            f.total_descuento, f.valor_iva, f.propina, f.importe_total,
            f.clave_acceso,
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=fila, column=col, value=v)
            c.fill = fill; c.font = Font(size=9)
        _fila_moneda(ws, fila, MONED_FC)

    if facturas:
        _fila_total(ws, len(facturas) + 2, 2, MONED_FC, COLOR_TOTAL)
    logger.info(f"FC_Recibidas: {len(facturas)} facturas")


def escribir_retenciones(ws, retenciones: list):
    ws.title = "CR_Recibidas"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 35
    _encabezados(ws, COLS_CR, COLOR_H_CR)

    TIPO_IMP = {"1": "RENTA", "2": "IVA", "6": "ISD"}
    fila = 2
    for r in retenciones:
        numero = f"{r.establecimiento}-{r.punto_emision}-{r.secuencial}"
        for imp in (r.impuestos or [None]):
            fill = PatternFill("solid", fgColor=COLOR_PAR if fila % 2 == 0 else COLOR_IMPAR)
            vals = [
                r.fecha_emision, r.ruc, r.razon_social, numero,
                r.identificacion_sujeto, r.razon_social_sujeto, r.periodo_fiscal,
                getattr(imp, "cod_doc_sustento", ""),
                getattr(imp, "num_doc_sustento", ""),
                getattr(imp, "fecha_emision_doc_sustento", ""),
                TIPO_IMP.get(getattr(imp, "codigo", ""), getattr(imp, "codigo", "")),
                getattr(imp, "codigo_retencion", ""),
                getattr(imp, "base_imponible", 0),
                getattr(imp, "porcentaje_retener", 0),
                getattr(imp, "valor_retenido", 0),
                r.clave_acceso,
            ]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=fila, column=col, value=v)
                c.fill = fill; c.font = Font(size=9)
            _fila_moneda(ws, fila, MONED_CR)
            fila += 1

    if fila > 2:
        _fila_total(ws, fila, 2, MONED_CR, COLOR_TOTAL)
    logger.info(f"CR_Recibidas: {len(retenciones)} retenciones")


# ── Hoja "Faltantes": comprobantes del TXT que no se pudieron descargar ──────
COLOR_H_FALT = "C00000"   # rojo

COLS_FALT = [
    ("CLAVE ACCESO", 52), ("TIPO", 14), ("RUC EMISOR", 15),
    ("RAZON SOCIAL", 35), ("NUMERO", 18), ("FECHA EMISION", 14),
    ("IMPORTE TOTAL", 14), ("MOTIVO", 55),
]


def escribir_faltantes(ws, faltantes: list):
    ws.title = "Faltantes"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    _encabezados(ws, COLS_FALT, COLOR_H_FALT)

    for fila, x in enumerate(faltantes, 2):
        fill = PatternFill("solid", fgColor=COLOR_PAR if fila % 2 == 0 else COLOR_IMPAR)
        vals = [
            x.get("clave", ""), x.get("tipo", ""), x.get("ruc_emisor", ""),
            x.get("razon", ""), x.get("numero", ""), x.get("fecha", ""),
            x.get("total", ""), x.get("motivo", ""),
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=fila, column=col, value=v)
            c.fill = fill; c.font = Font(size=9)
        ws.cell(row=fila, column=7).number_format = '#,##0.00'
    logger.info(f"Faltantes: {len(faltantes)}")


def generar_reporte(facturas: list, retenciones: list,
                    ruta_salida=None, abrir_al_finalizar: bool = False,
                    faltantes: list = None) -> Path:
    if ruta_salida is None:
        import config
        ruta_salida = config.OUTPUT_DIR / f"SRI_Reporte_{timestamp()}.xlsx"
    ruta_salida = Path(ruta_salida)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    escribir_facturas(wb.active, facturas)
    escribir_retenciones(wb.create_sheet(), retenciones)
    if faltantes:
        escribir_faltantes(wb.create_sheet(), faltantes)
    wb.save(ruta_salida)
    logger.info(f"Reporte guardado: {ruta_salida}")

    if abrir_al_finalizar:
        import os
        try:
            os.startfile(str(ruta_salida))
        except Exception:
            pass
    return ruta_salida


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE DE COMPROBANTES EMITIDOS (Versión 2)
#   Una hoja por tipo: FC_Emitidas, NC_Emitidas, LIQ_Emitidas, RET_Emitidas
# ══════════════════════════════════════════════════════════════════════════════

COLOR_H_EMIT = "1F6E43"   # verde para emitidas

# Columnas de comprobantes "tipo venta" emitidos (FC / NC / LIQ)
COLS_EMIT_FC = [
    ("FECHA EMISION", 16), ("FECHA AUTORIZACION", 20), ("RUC CONTRAPARTE", 15),
    ("CONTRAPARTE", 35), ("NUMERO", 18),
    ("BASE 0%", 12), ("BASE 12%", 12), ("BASE 14%", 12), ("BASE 15%", 12),
    ("BASE NO OBJETO", 14), ("BASE EXENTA", 12), ("SUBTOTAL", 12),
    ("DESCUENTO", 12), ("IVA", 12), ("TOTAL", 12), ("CLAVE ACCESO", 52),
]
MONED_EMIT_FC = [6, 7, 8, 9, 10, 11, 12, 13, 14, 15]

# Columnas del reporte de retenciones emitidas (formato pivoteado del modelo)
COLS_EMIT_RET = [
    ("EJERCICIO FISCAL", 16), ("SUJETO RETENIDO", 38), ("RUC SUJETO", 16),
    ("COMPROBANTE", 16), ("NUMERO", 20), ("FECHA EMISION", 14),
    ("FECHA AUTORIZACION", 20), ("SECUENCIAL", 20), ("CLAVE ACCESO", 52),
    ("BASE RENTA", 13), ("% RENTA", 10), ("VALOR RENTA", 13),
    ("BASE IVA", 13), ("% IVA", 10), ("VALOR IVA", 13), ("TOTAL", 13),
    # Codigos al final, para no alterar el orden del modelo original
    ("COD. RENTA", 12), ("COD. IVA", 12),
]
MONED_EMIT_RET = [10, 12, 13, 15, 16]   # bases, valores y total (no los %)
PCT_EMIT_RET   = [11, 14]               # % renta, % iva

_DOC_SUSTENTO_NOM = {
    "01": "FACTURA", "02": "N. VENTA", "03": "LIQ DE COMPRAS",
    "04": "N. CREDITO", "05": "N. DEBITO", "06": "GUIA REMISION",
    "07": "RETENCION", "11": "PASAJE", "12": "DOC FINANCIERO",
    "19": "COMPROB PAGO", "20": "RECEPCION", "21": "CARTA PORTE",
}


def _escribir_emit_facturas_like(ws, items: list, titulo: str, contraparte_lbl: str):
    """Escribe una hoja de comprobantes tipo venta emitidos (FC/NC/LIQ)."""
    ws.title = titulo
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    cols = list(COLS_EMIT_FC)
    cols[2] = (f"RUC {contraparte_lbl}", 15)
    cols[3] = (contraparte_lbl, 35)
    _encabezados(ws, cols, COLOR_H_EMIT)

    for fila, f in enumerate(items, 2):
        fill = PatternFill("solid", fgColor=COLOR_PAR if fila % 2 == 0 else COLOR_IMPAR)
        vals = [
            f.fecha_emision, f.fecha_autorizacion,
            f.identificacion_comprador, f.razon_social_comprador, f.numero,
            f.base_iva_0, f.base_iva_12, f.base_iva_14, f.base_iva_15,
            f.base_no_objeto, f.base_exenta, f.total_sin_impuestos,
            f.total_descuento, f.valor_iva, f.importe_total, f.clave_acceso,
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=fila, column=col, value=v)
            c.fill = fill; c.font = Font(size=9)
        _fila_moneda(ws, fila, MONED_EMIT_FC)

    if items:
        _fila_total(ws, len(items) + 2, 2, MONED_EMIT_FC, COLOR_TOTAL)
    logger.info(f"{titulo}: {len(items)} comprobantes")


def _escribir_emit_retenciones(ws, retenciones: list):
    """
    Hoja RET_Emitidas con el formato pivoteado del modelo:
    una fila por documento de sustento, con Renta e IVA en columnas separadas.

    Cuando un mismo documento de sustento trae VARIAS retenciones de renta o de
    IVA (porcentajes distintos sobre la misma factura), se emite una fila
    adicional por cada una para no perder ningun valor. Se emparejan por
    posicion: la fila N lleva la N-esima de renta junto a la N-esima de IVA.
    Ese emparejamiento es solo un acomodo visual — cada celda conserva su
    propia base, porcentaje, valor y codigo.
    """
    ws.title = "RET_Emitidas"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    _encabezados(ws, COLS_EMIT_RET, COLOR_H_EMIT)

    fila = 2
    for r in retenciones:
        # Agrupar impuestos por documento de sustento (codDoc + numDoc)
        grupos = {}
        for imp in r.impuestos:
            clave = (imp.cod_doc_sustento, imp.num_doc_sustento)
            grupos.setdefault(clave, []).append(imp)
        if not grupos:
            grupos = {("", ""): []}

        escrito = 0.0
        for (cod_sus, num_sus), imps in grupos.items():
            rentas = [i for i in imps if i.codigo == "1"]
            ivas   = [i for i in imps if i.codigo == "2"]

            for k in range(max(len(rentas), len(ivas), 1)):
                renta = rentas[k] if k < len(rentas) else None
                iva   = ivas[k]   if k < len(ivas)   else None
                fill = PatternFill("solid",
                                   fgColor=COLOR_PAR if fila % 2 == 0 else COLOR_IMPAR)

                v_renta = renta.valor_retenido if renta else 0
                v_iva   = iva.valor_retenido   if iva   else 0
                escrito += v_renta + v_iva

                vals = [
                    r.periodo_fiscal,
                    r.razon_social_sujeto,
                    r.identificacion_sujeto,
                    _DOC_SUSTENTO_NOM.get(cod_sus, cod_sus),
                    num_sus,
                    r.fecha_emision,
                    r.fecha_autorizacion,
                    r.numero,
                    r.clave_acceso,
                    renta.base_imponible if renta else 0,
                    (renta.porcentaje_retener / 100.0) if renta else 0,
                    v_renta,
                    iva.base_imponible if iva else 0,
                    (iva.porcentaje_retener / 100.0) if iva else 0,
                    v_iva,
                    v_renta + v_iva,
                    renta.codigo_retencion if renta else "",
                    iva.codigo_retencion if iva else "",
                ]
                for col, v in enumerate(vals, 1):
                    c = ws.cell(row=fila, column=col, value=v)
                    c.fill = fill; c.font = Font(size=9)
                _fila_moneda(ws, fila, MONED_EMIT_RET)
                for col in PCT_EMIT_RET:
                    ws.cell(row=fila, column=col).number_format = '0.0%'
                fila += 1

        # Red de seguridad: lo escrito debe cuadrar con el total del XML.
        # Si no cuadra, queda registrado en el log en vez de pasar silencioso.
        if abs(escrito - r.total_retenido) > 0.01:
            logger.warning(
                f"RET {r.numero}: las filas suman {escrito:.2f} pero el XML "
                f"declara {r.total_retenido:.2f} — revisar")

    if fila > 2:
        # Total solo en columnas monetarias
        _fila_total(ws, fila, 2, MONED_EMIT_RET, COLOR_TOTAL)
    logger.info(f"RET_Emitidas: {len(retenciones)} retenciones, {fila - 2} filas")


def generar_reporte_emitidos(facturas: list, retenciones: list,
                             ruta_salida=None, abrir_al_finalizar: bool = False,
                             faltantes: list = None) -> Path:
    """
    Genera el reporte de comprobantes EMITIDOS con una hoja por tipo.
    'facturas' incluye FC/NC/ND/LIQ (separadas aquí por cod_doc).
    """
    if ruta_salida is None:
        import config
        ruta_salida = config.OUTPUT_DIR / f"SRI_Emitidos_{timestamp()}.xlsx"
    ruta_salida = Path(ruta_salida)
    ruta_salida.parent.mkdir(parents=True, exist_ok=True)

    fc  = [f for f in facturas if getattr(f, "cod_doc", "01") == "01"]
    nc  = [f for f in facturas if getattr(f, "cod_doc", "") in ("04", "05")]
    liq = [f for f in facturas if getattr(f, "cod_doc", "") == "03"]

    wb = Workbook()
    _escribir_emit_facturas_like(wb.active,            fc,  "FC_Emitidas",  "CLIENTE")
    _escribir_emit_facturas_like(wb.create_sheet(),    nc,  "NC_Emitidas",  "CLIENTE")
    _escribir_emit_facturas_like(wb.create_sheet(),    liq, "LIQ_Emitidas", "PROVEEDOR")
    _escribir_emit_retenciones(wb.create_sheet(),      retenciones)
    if faltantes:
        escribir_faltantes(wb.create_sheet(), faltantes)
    wb.save(ruta_salida)
    logger.info(f"Reporte emitidos guardado: {ruta_salida}")

    if abrir_al_finalizar:
        import os
        try:
            os.startfile(str(ruta_salida))
        except Exception:
            pass
    return ruta_salida
